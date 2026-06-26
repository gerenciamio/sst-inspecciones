from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
import sqlite3, os, io, json, base64
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from PIL import Image as PILImage

app = Flask(__name__)
CORS(app)

DB_PATH = os.environ.get('DB_PATH', 'inspecciones.db')

# ── Database ──────────────────────────────────────────────────────────────────
def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    with get_db() as conn:
        conn.execute('''
            CREATE TABLE IF NOT EXISTS visitas (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                cliente TEXT NOT NULL,
                fecha TEXT NOT NULL,
                participantes TEXT,
                asesora TEXT,
                creado TEXT NOT NULL
            )
        ''')
        conn.execute('''
            CREATE TABLE IF NOT EXISTS hallazgos (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                visita_id INTEGER NOT NULL,
                lugar TEXT,
                situacion TEXT,
                recomendacion TEXT,
                foto_antes TEXT,
                foto_despues TEXT,
                estado TEXT DEFAULT 'pendiente',
                FOREIGN KEY(visita_id) REFERENCES visitas(id)
            )
        ''')
        conn.commit()

init_db()

# ── Static files ──────────────────────────────────────────────────────────────
@app.route('/')
def index():
    with open('index.html', encoding='utf-8') as f:
        return f.read()

@app.route('/admin')
def admin():
    with open('admin.html', encoding='utf-8') as f:
        return f.read()

@app.route('/cliente')
def cliente():
    with open('cliente.html', encoding='utf-8') as f:
        return f.read()

# ── API: create visit + hallazgos ─────────────────────────────────────────────
@app.route('/api/visita', methods=['POST'])
def crear_visita():
    data = request.get_json()
    cliente = data.get('cliente','').strip()
    fecha = data.get('fecha','').strip()
    if not cliente or not fecha:
        return jsonify({'error': 'Cliente y fecha requeridos'}), 400
    ahora = datetime.now().strftime('%Y-%m-%d %H:%M')
    with get_db() as conn:
        cur = conn.execute(
            'INSERT INTO visitas (cliente, fecha, participantes, asesora, creado) VALUES (?,?,?,?,?)',
            (cliente, fecha, data.get('participantes',''), data.get('asesora',''), ahora)
        )
        visita_id = cur.lastrowid
        for h in data.get('hallazgos', []):
            conn.execute(
                'INSERT INTO hallazgos (visita_id, lugar, situacion, recomendacion, foto_antes, estado) VALUES (?,?,?,?,?,?)',
                (visita_id, h.get('lugar',''), h.get('situacion',''), h.get('recomendacion',''), h.get('fotoBefore',''), 'pendiente')
            )
        conn.commit()
    return jsonify({'ok': True, 'visita_id': visita_id})

# ── API: add hallazgo to existing visit ───────────────────────────────────────
@app.route('/api/visita/<int:visita_id>/hallazgo', methods=['POST'])
def agregar_hallazgo(visita_id):
    data = request.get_json()
    with get_db() as conn:
        conn.execute(
            'INSERT INTO hallazgos (visita_id, lugar, situacion, recomendacion, foto_antes, estado) VALUES (?,?,?,?,?,?)',
            (visita_id, data.get('lugar',''), data.get('situacion',''), data.get('recomendacion',''), data.get('fotoBefore',''), 'pendiente')
        )
        conn.commit()
    return jsonify({'ok': True})

# ── API: list visits ──────────────────────────────────────────────────────────
@app.route('/api/visitas')
def listar_visitas():
    cliente_q = request.args.get('cliente', '')
    fecha_q = request.args.get('fecha', '')
    with get_db() as conn:
        query = '''SELECT v.*, 
            COUNT(h.id) as total_hallazgos,
            SUM(CASE WHEN h.estado='cerrado' THEN 1 ELSE 0 END) as cerrados
            FROM visitas v LEFT JOIN hallazgos h ON h.visita_id=v.id
            WHERE 1=1'''
        params = []
        if cliente_q:
            query += ' AND LOWER(v.cliente) LIKE ?'
            params.append(f'%{cliente_q.lower()}%')
        if fecha_q:
            query += ' AND v.fecha=?'
            params.append(fecha_q)
        query += ' GROUP BY v.id ORDER BY v.creado DESC LIMIT 100'
        rows = conn.execute(query, params).fetchall()
    return jsonify([dict(r) for r in rows])

# ── API: get visit detail with hallazgos ──────────────────────────────────────
@app.route('/api/visita/<int:visita_id>')
def detalle_visita(visita_id):
    with get_db() as conn:
        v = conn.execute('SELECT * FROM visitas WHERE id=?', (visita_id,)).fetchone()
        if not v:
            return jsonify({'error': 'No encontrada'}), 404
        hs = conn.execute('SELECT * FROM hallazgos WHERE visita_id=? ORDER BY id', (visita_id,)).fetchall()
    return jsonify({'visita': dict(v), 'hallazgos': [dict(h) for h in hs]})

# ── API: upload after photo ───────────────────────────────────────────────────
@app.route('/api/hallazgo/<int:hallazgo_id>/despues', methods=['POST'])
def subir_despues(hallazgo_id):
    data = request.get_json()
    foto = data.get('foto_despues', '')
    with get_db() as conn:
        conn.execute(
            "UPDATE hallazgos SET foto_despues=?, estado='cerrado' WHERE id=?",
            (foto, hallazgo_id)
        )
        conn.commit()
    return jsonify({'ok': True})

# ── API: export Excel ─────────────────────────────────────────────────────────
def b64_to_xl_image(b64_str, max_w=200, max_h=150):
    try:
        data = base64.b64decode(b64_str.split(',')[1])
        img = PILImage.open(io.BytesIO(data)).convert('RGB')
        img.thumbnail((max_w, max_h), PILImage.LANCZOS)
        buf = io.BytesIO()
        img.save(buf, format='PNG')
        buf.seek(0)
        return XLImage(buf)
    except:
        return None

@app.route('/api/exportar/<int:visita_id>')
def exportar_excel(visita_id):
    with get_db() as conn:
        v = conn.execute('SELECT * FROM visitas WHERE id=?', (visita_id,)).fetchone()
        if not v:
            return jsonify({'error': 'No encontrada'}), 404
        hs = conn.execute('SELECT * FROM hallazgos WHERE visita_id=? ORDER BY id', (visita_id,)).fetchall()

    v = dict(v)
    hs = [dict(h) for h in hs]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Informe de inspección'

    thin = Side(style='thin', color='000000')
    ALL_BORDERS = Border(top=thin, bottom=thin, left=thin, right=thin)

    def apply(cell, val, bold=False, sz=10, h='center', v_align='center', fill=None, border=False):
        cell.value = val
        cell.font = Font(name='Arial', bold=bold, size=sz)
        cell.alignment = Alignment(horizontal=h, vertical=v_align, wrap_text=True)
        if fill:
            cell.fill = PatternFill('solid', fgColor=fill)
        if border:
            cell.border = ALL_BORDERS

    # Column widths
    for col, w in [('A',16),('B',10),('C',12),('D',21),('E',38),('F',38),('G',25),('H',12)]:
        ws.column_dimensions[col].width = w

    # Row heights header
    for i in range(1,6):
        ws.row_dimensions[i].height = [14,14,33,33,35][i-1]

    # Header
    apply(ws['A1'], 'LOGO', bold=True)
    apply(ws['C1'], 'SEGURIDAD SALUD EN EL TRABAJO', bold=True)
    apply(ws['F1'], f'Fecha: {v["fecha"]}', bold=True, sz=8)
    apply(ws['C2'], 'INFORME Y SEGUIMIENTO A INSPECCIONES', bold=True)
    apply(ws['F2'], 'CODIGO: FT-SST-020', bold=True, sz=8)
    apply(ws['F3'], 'Pag 1 de 2', bold=True, sz=8)
    apply(ws['G3'], 'Versión 1', bold=True)
    apply(ws['A4'], f'Fecha inspección: {v["fecha"]}', bold=True, h='left')
    apply(ws['E4'], f'PARTICIPANTES: {v["participantes"]}', bold=True, h='left')

    gray = 'D3D1C7'
    for col, txt in [('A','LUGAR'),('B','SITUACIÓN ENCONTRADA'),('E','FOTO ANTES'),('F','FOTO DESPUÉS'),('G','RECOMENDACIÓN'),('H','ESTADO')]:
        apply(ws[f'{col}5'], txt, bold=True, fill=gray, border=True)

    # Merges header
    for merge in ['A1:B3','C1:E1','C2:E3','F1:G1','F2:G2','F3:G3','A4:B4','E4:H4','B5:D5']:
        ws.merge_cells(merge)

    ROW_H = 140
    for i, h in enumerate(hs):
        r = 6 + i
        ws.row_dimensions[r].height = ROW_H
        apply(ws[f'A{r}'], h.get('lugar',''), border=True, h='left', v_align='top')
        apply(ws[f'B{r}'], h.get('situacion',''), border=True, h='left', v_align='top')
        apply(ws[f'G{r}'], h.get('recomendacion',''), border=True, h='left', v_align='top')
        estado = h.get('estado','pendiente')
        apply(ws[f'H{r}'], 'Cerrado' if estado=='cerrado' else 'Pendiente',
              border=True, fill='C0DD97' if estado=='cerrado' else 'FAC775')
        apply(ws[f'E{r}'], '', border=True)
        apply(ws[f'F{r}'], '', border=True)
        ws.merge_cells(f'B{r}:D{r}')

        if h.get('foto_antes'):
            img = b64_to_xl_image(h['foto_antes'])
            if img:
                img.anchor = f'E{r}'
                ws.add_image(img)

        if h.get('foto_despues'):
            img2 = b64_to_xl_image(h['foto_despues'])
            if img2:
                img2.anchor = f'F{r}'
                ws.add_image(img2)

    obs_row = 6 + len(hs)
    sig_row = obs_row + 1
    ws.row_dimensions[obs_row].height = 80
    ws.row_dimensions[sig_row].height = 50
    apply(ws[f'A{obs_row}'], 'OBSERVACIONES GENERALES', h='left', v_align='top')
    ws.merge_cells(f'A{obs_row}:H{obs_row}')
    apply(ws[f'A{sig_row}'], 'Firma de quien realizó la inspección' + ' '*35 + 'Firma de quien recibió la inspección', sz=11, h='center')
    ws.merge_cells(f'A{sig_row}:H{sig_row}')

    ws['!ref'] = None
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"FT-SST-020_{v['cliente'].replace(' ','_')}_{v['fecha']}.xlsx"
    return send_file(buf, as_attachment=True, download_name=fname,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port)
